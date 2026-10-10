import io
import re
from typing import List, Dict, Any
import pypdf
import docx
from .sensitive_detector import PATTERNS, CATEGORY_ROLE_MAP

SECRET_PATTERNS = {
    "API Key / Bearer Token": re.compile(
        r"(?i)(?:api_key|apikey|secret|token|bearer|auth_token)\s*[:=]\s*['\"]?([a-zA-Z0-9_\-]{16,})['\"]?"
    ),
    "Private Key Block": re.compile(
        r"-----BEGIN (?:RSA |EC )?PRIVATE KEY-----"
    ),
    "Database Connection String": re.compile(
        r"(?i)(?:postgres|postgresql|mysql|mongodb|redis):\/\/[a-zA-Z0-9_\-\.]+:[^@\s]+@[a-zA-Z0-9_\-\.]+"
    ),
    "AWS Access Key": re.compile(
        r"\b(AKIA[0-9A-Z]{16})\b"
    ),
}

def mask_evidence(text: str, keep_chars: int = 4) -> str:
    """Masks secret or sensitive snippet in report so raw secrets aren't leaked in findings."""
    clean = text.strip()
    if len(clean) <= keep_chars * 2:
        return "*" * len(clean)
    return clean[:keep_chars] + "*" * (len(clean) - keep_chars * 2) + clean[-keep_chars:]

def extract_text_from_pdf(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Extracts text page by page from PDF bytes."""
    pages_content = []
    try:
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        for page_idx, page in enumerate(reader.pages):
            text = page.extract_text() or ""
            pages_content.append({
                "page": page_idx + 1,
                "text": text
            })
    except Exception as e:
        pages_content.append({"page": 1, "text": f"[Error reading PDF: {str(e)}]"})
    return pages_content

def extract_text_from_docx(file_bytes: bytes) -> List[Dict[str, Any]]:
    """Extracts text paragraph by paragraph and tables from DOCX."""
    doc = docx.Document(io.BytesIO(file_bytes))
    sections = []
    
    # Paragraphs
    for p_idx, p in enumerate(doc.paragraphs):
        if p.text.strip():
            sections.append({
                "location": f"Paragraph {p_idx + 1}",
                "text": p.text
            })
            
    # Tables
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                sections.append({
                    "location": f"Table {t_idx + 1}, Row {r_idx + 1}",
                    "text": " | ".join(row_texts)
                })
    return sections

def scan_text_content(content_blocks: List[Dict[str, Any]], source_name: str) -> List[Dict[str, Any]]:
    """
    Scans extracted text blocks across lines/paragraphs/pages for PII and secrets.
    """
    findings = []

    for block in content_blocks:
        location = block.get("location") or f"Page {block.get('page', 1)}"
        text = block.get("text", "")

        lines = text.split("\n")
        for line_num, line in enumerate(lines, 1):
            stripped_line = line.strip()
            if not stripped_line:
                continue

            # 1. PII Patterns
            for pat_name, pat_regex in PATTERNS.items():
                for match in pat_regex.finditer(stripped_line):
                    matched_val = match.group(0)
                    findings.append({
                        "Source": source_name,
                        "Location": f"{location}, Line {line_num}",
                        "Category": pat_name,
                        "Role": CATEGORY_ROLE_MAP.get(pat_name, "Direct Identifier"),
                        "Detection Method": "Pattern Regex",
                        "Evidence": mask_evidence(matched_val),
                        "Confidence": "High",
                        "Recommended Action": "Redact or remove from document before sharing."
                    })

            # 2. Secret & Credential Patterns
            for secret_type, secret_regex in SECRET_PATTERNS.items():
                for match in secret_regex.finditer(stripped_line):
                    matched_val = match.group(0)
                    findings.append({
                        "Source": source_name,
                        "Location": f"{location}, Line {line_num}",
                        "Category": "Authentication Secret",
                        "Role": "Direct Secret",
                        "Detection Method": secret_type,
                        "Evidence": mask_evidence(matched_val, keep_chars=3),
                        "Confidence": "Very High",
                        "Recommended Action": "Rotate credential immediately and remove from file."
                    })

    return findings

def scan_document(file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    """Entrypoint to scan any supported document/code/config file."""
    lower_name = filename.lower()
    
    if lower_name.endswith(".pdf"):
        blocks = extract_text_from_pdf(file_bytes)
        return scan_text_content(blocks, filename)
    elif lower_name.endswith(".docx"):
        blocks = extract_text_from_docx(file_bytes)
        return scan_text_content(blocks, filename)
    else:
        # Plain text, markdown, json, python, yaml, env, logs
        try:
            text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            text = file_bytes.decode("latin-1", errors="ignore")

        blocks = [{"location": f"Line {i+1}", "text": line} for i, line in enumerate(text.splitlines())]
        return scan_text_content(blocks, filename)
